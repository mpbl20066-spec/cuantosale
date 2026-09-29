# -*- coding: utf-8 -*-
"""Carga tours.xlsx (salida del scraper) a public.tours de Supabase.

    python scripts/cargar-tours-scraper.py --dry-run          # solo informa
    python scripts/cargar-tours-scraper.py --json filas.json  # escribe las filas para revisar
    python scripts/cargar-tours-scraper.py                    # sube

Columnas de la hoja "Tours": region, locality, country, source, agency, title,
category, price, currency, duration_min, rating, reviews, languages,
free_cancellation, url, description, scraped_at.

Usa el mismo RPC que cargar-tours-xlsx.py (tours_guardar_lote), asi que la
tabla sigue teniendo un solo camino de escritura.

Reglas:
- El destino sale de "locality" (no de "region", que agrupa varias ciudades).
  Cada locality esta mapeada a mano a la key de lib/model.js. Una locality
  sin mapa se avisa y no se sube: mandarla a una key inventada deja el tour
  invisible sin error.
- Precio en USD -> columna precio (la que muestra la web). Precio en BRL ->
  precio_brl (origen; el server lo convierte). Otra moneda o sin precio: no
  se sube (un tour activo necesita precio) y va a data/tours-sin-precio.json.
- Este archivo no trae transfers ni privados: todo lo que trae es tour.
- Duplicados (mismo destino y titulo): gana el de menor precio.
"""
import os
import re
import sys
import json
import argparse
import unicodedata

from openpyxl import load_workbook

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(os.environ.get("USERPROFILE", ""), "Desktop", "tours.xlsx")

LOCALITY = {
    "arraial d'ajuda": "ajuda", "itacare": "itacare", "morro de sao paulo": "morro",
    "porto seguro": "portoseguro", "praia do forte": "forte", "salvador de bahia": "ssa",
    "trancoso": "trancoso", "belo horizonte": "bho", "buenos aires": "bue",
    "arraial do cabo": "arraial", "buzios": "buz", "cabo frio": "cabo",
    "angra dos reis": "angra", "ilha grande": "ilha", "ilhabela": "ilhabela",
    "paraty": "paraty", "ubatuba": "ubatuba", "curitiba": "curitiba",
    "foz do iguacu": "igu", "jericoacoara": "jericoacoara",
    "capao da canoa": "canoa", "torres": "torres",
    "balneario camboriu": "bcm", "bombinhas": "bombinhas", "florianopolis": "fln",
    "garopaba": "garopaba", "picarras": "picarras", "praia do rosa": "rosa",
    "fernando de noronha": "fernando", "fortaleza": "for", "joao pessoa": "joaopessoa",
    "maceio": "mcz", "maragogi": "maragogi", "natal": "nat", "pipa": "pip",
    "porto de galinhas": "porto", "recife": "rec", "canela": "canela",
    "gramado": "gram", "porto alegre": "poa", "sao paulo": "sao",
    "rio de janeiro": "rio",
}


def cargar_env():
    try:
        with open(os.path.join(BASE, ".env"), encoding="utf-8") as f:
            for linea in f:
                m = re.match(r"^\s*([A-Z0-9_]+)\s*=\s*(.*?)\s*$", linea.rstrip("\n"))
                if m and m.group(1) not in os.environ:
                    os.environ[m.group(1)] = m.group(2).strip("'\"")
    except (IOError, OSError):
        pass


def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return " ".join(s.lower().split())


def texto(v):
    return " ".join(str(v).split()) if v is not None else ""


def duracion(minutos):
    if not minutos:
        return ""
    m = int(round(float(minutos)))
    h, r = divmod(m, 60)
    if h and r:
        return "%d h %d min" % (h, r)
    return "%d h" % h if h else "%d min" % r


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--archivo", default=XLSX)
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--json")
    args = ap.parse_args()

    cargar_env()
    if not os.path.exists(args.archivo):
        sys.exit("No existe: " + args.archivo)
    ws = load_workbook(args.archivo, data_only=True, read_only=True)["Tours"]
    it = ws.iter_rows(values_only=True)
    head = [norm(h) for h in next(it)]
    col = {h: i for i, h in enumerate(head)}
    faltan = [c for c in ("locality", "title", "price", "currency", "url") if c not in col]
    if faltan:
        sys.exit("Faltan columnas en la hoja Tours: " + ", ".join(faltan))

    def g(f, k):
        i = col.get(k)
        return f[i] if i is not None and i < len(f) else None

    por_clave = {}
    sin_precio, sin_destino = [], {}
    total = 0
    for n, f in enumerate(it, start=2):
        titulo = texto(g(f, "title"))
        if not titulo:
            continue
        total += 1
        loc = texto(g(f, "locality"))
        destino = LOCALITY.get(norm(loc))
        if not destino:
            sin_destino[loc] = sin_destino.get(loc, 0) + 1
            continue
        try:
            precio_v = float(g(f, "price")) if g(f, "price") not in (None, "") else None
        except (TypeError, ValueError):
            precio_v = None
        moneda = norm(g(f, "currency"))
        if not precio_v or moneda not in ("usd", "brl"):
            sin_precio.append({"destino": loc, "titulo": titulo, "moneda": moneda,
                               "agencia": texto(g(f, "agency")), "hoja": "Tours", "fila": n})
            continue
        fila = {
            "destino": destino,
            "titulo": titulo,
            "descripcion": texto(g(f, "description")),
            "detalle": "",
            "precio": precio_v if moneda == "usd" else None,
            "precio_brl": precio_v if moneda == "brl" else None,
            "pvp": precio_v,
            "duracion": duracion(g(f, "duration_min")),
            "tipo_servicio": texto(g(f, "category")) or "Tour",
            "politica_cancelacion": "Cancelación gratuita" if g(f, "free_cancellation") else "",
            "link_web": texto(g(f, "url")),
            "agencia": texto(g(f, "agency")) or texto(g(f, "source")),
            "estado_scrapeo": "ok",
            "activo": True,
            "fuente": "Scraper %s (%s), %s" % (texto(g(f, "source")), os.path.basename(args.archivo),
                                               texto(g(f, "scraped_at"))[:10]),
            "verificado": None,
        }
        k = (destino, titulo)
        previo = por_clave.get(k)
        if previo is None or precio_v < (previo["pvp"] or 1e12):
            por_clave[k] = fila

    filas = list(por_clave.values())
    cont = {}
    for f in filas:
        cont[f["destino"]] = cont.get(f["destino"], 0) + 1
        f["orden"] = cont[f["destino"]] * 10

    print("de %d filas con titulo:" % total)
    print("  %d sin precio o con moneda no soportada (data/tours-sin-precio.json)" % len(sin_precio))
    print("  %d con locality sin mapa: %s" % (sum(sin_destino.values()), sin_destino or "-"))
    print("  %d duplicados fusionados" % (total - len(sin_precio) - sum(sin_destino.values()) - len(filas)))
    print("a subir: %d tours en %d destinos" % (len(filas), len(cont)))
    for k in sorted(cont):
        print("    %-14s %d" % (k, cont[k]))

    if sin_precio:
        with open(os.path.join(BASE, "data", "tours-sin-precio.json"), "w", encoding="utf-8") as fh:
            json.dump(sin_precio, fh, ensure_ascii=False, indent=2)
    if args.json:
        with open(args.json, "w", encoding="utf-8") as fh:
            json.dump(filas, fh, ensure_ascii=False, indent=2)
        print("escrito:", args.json)
        return
    if args.dry_run:
        print("--dry-run: no se escribio nada en Supabase.")
        return
    if not filas:
        sys.exit("No hay filas para subir.")

    url = os.environ.get("SUPABASE_URL", "").strip().rstrip("/")
    key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY", "").strip()
    if not url or not key:
        sys.exit("Faltan SUPABASE_URL o SUPABASE_SERVICE_ROLE_KEY (la anon no sirve: la tabla esta cerrada).")

    import urllib.request
    import urllib.error
    for i in range(0, len(filas), 100):
        lote = filas[i:i + 100]
        req = urllib.request.Request(
            url + "/rest/v1/rpc/tours_guardar_lote",
            data=json.dumps({"p_filas": lote}).encode("utf-8"),
            headers={"apikey": key, "Authorization": "Bearer " + key,
                     "Content-Type": "application/json", "Prefer": "return=minimal"},
            method="POST")
        try:
            urllib.request.urlopen(req, timeout=60).read()
        except urllib.error.HTTPError as e:
            sys.exit("PostgREST %s en el lote %d: %s" % (e.code, i // 100 + 1, e.read().decode("utf-8", "replace")[:600]))
        print("  %d/%d" % (min(i + 100, len(filas)), len(filas)))
    print("cargados %d tours en public.tours" % len(filas))


if __name__ == "__main__":
    main()
