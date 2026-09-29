# -*- coding: utf-8 -*-
"""Lee los costos de data/tabla-transfers.xlsx y genera el SQL para subirlos.

    python scripts/build-costos-desde-xlsx.py

Escribe supabase_costos_cargados.sql, que se pega en el SQL Editor de Supabase.
Solo emite filas que tienen ALGO cargado: los 45 destinos ya los crea
supabase_transfer_costos.sql, asi que no hay que volver a mandarlos.

Este script es el puente para trabajar en Excel. La fuente de verdad del costo
es la tabla transfer_destinos: el .xlsx se regenera con `npm run db:pull` y lo
que se escriba a mano en el proximo pull se pierde. Es un ida y vuelta, no una
copia permanente.
"""
import os, sys, unicodedata
from openpyxl import load_workbook

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(BASE, "data", "tabla-transfers.xlsx")
JSON = os.path.join(BASE, "data", "transfer-precios.json")
OUT = os.path.join(BASE, "supabase_costos_cargados.sql")

import json as _json


def norm(s):
    """'Búzios' y 'Buzios' tienen que ser el mismo destino."""
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return " ".join(s.lower().split())


if not os.path.exists(XLSX):
    sys.exit("No existe " + XLSX + ". Corré antes: python scripts/armar-tabla-transfers.py")

with open(JSON, encoding="utf-8") as f:
    por_nombre = {norm(v["nombre"]): k for k, v in _json.load(f)["destinos"].items()}

ws = load_workbook(XLSX, data_only=False)["Transfers"]
head = [c.value for c in ws[1]]
try:
    COL = {h: head.index(h) + 1 for h in
           ("Destino", "Costo Base Compartido", "Costo Base Privado", "Comision")}
except ValueError:
    sys.exit("La fila 1 no tiene las columnas esperadas. Regenerá el .xlsx "
             "con scripts/armar-tabla-transfers.py")


def numero(v):
    """Acepta 88, '88', '$ 88', 'US$ 88', '30%', '1.200' (miles). Devuelve float o None.

    El separador de miles importa porque el costo de un transfer privado puede
    pasar de mil: escrito a mano en una celda de texto, '1.200' son mil doscientos
    en formato es y 1,2 en formato en. Se resuelve mirando cuantos digitos van
    despues del punto: tres son miles, otra cosa es decimal.
    """
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s:
        return None
    s = s.replace("%", "").replace("US$", "").replace("$", "").replace("R$", "")
    s = s.replace("\u00a0", "").replace(" ", "").strip()
    if "," in s:
        # Con coma, el punto es de miles y la coma es el decimal.
        s = s.replace(".", "").replace(",", ".")
    elif "." in s:
        partes = s.split(".")
        decimales = partes[-1]
        # Todos los puntos son de miles si el ultimo grupo tiene 3 digitos y hay
        # mas de un punto, o si hay uno solo seguido de exactamente 3 digitos.
        if len(decimales) == 3 and len(partes) >= 2:
            s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return None


filas, avisos = [], []
for r in range(2, 47):
    nombre = ws.cell(row=r, column=COL["Destino"]).value
    if not nombre:
        continue
    cco = numero(ws.cell(row=r, column=COL["Costo Base Compartido"]).value)
    cpr = numero(ws.cell(row=r, column=COL["Costo Base Privado"]).value)
    com = numero(ws.cell(row=r, column=COL["Comision"]).value)

    if cco is None and cpr is None and com is None:
        continue

    # El destino tiene que existir: una fila con el nombre mal escrito se
    # pegaria a un destino que no esta en DEST y el costo quedariaorphan.
    clave = por_nombre.get(norm(nombre))
    if not clave:
        avisos.append("fila %d: '%s' no coincide con ningun destino de DEST. Se saltea."
                      % (r, nombre))
        continue

    if com is not None and com > 1:
        avisos.append("%s: la comision esta como %s. Es una fraccion: escribi 0,30 "
                      "o usa el formato 30%%. Se toma como %s."
                      % (nombre, com, round(com / 100, 4)))
        com = round(com / 100, 4)

    for campo, valor in (("costo compartido", cco), ("costo privado", cpr)):
        if valor is not None and valor < 0:
            avisos.append("%s: %s negativo (%s). Se descarta." % (nombre, campo, valor))
            if campo == "costo compartido":
                cco = None
            else:
                cpr = None

    filas.append((clave, nombre, cco, cpr, com))

# --- avisos ---
for a in avisos:
    print("aviso: " + a)
print("")

if not filas:
    print("No hay ningun costo cargado en el .xlsx. No se genero SQL.")
    print("El archivo de salida anterior, si habia, queda como estaba.")
    sys.exit(0)


def lit(v):
    return "null" if v is None else str(v)


valores = ",\n".join(
    "  ('%s', %s, %s, %s)   -- %s" % (k, lit(c), lit(p), lit(m), n)
    for k, n, c, p, m in filas
)

sql = """-- Costos de transfer cargados desde data/tabla-transfers.xlsx.
--
-- Generado por scripts/build-costos-desde-xlsx.py. Pegar en el SQL Editor de
-- Supabase y darle Run.
--
-- SOLO va lo que tiene costo cargado. Los 45 destinos y sus columnas de
-- referencia los crea supabase_transfer_costos.sql: si ese todavia no se corrio,
-- esto agrega las filas con on conflict, asi que tambien sirve para crearlas de
-- a poco.
--
-- comision va como fraccion (0.30 es 30%%) y los precios finales los calcula la
-- columna generada de la tabla, no esta consulta.

insert into public.transfer_destinos
  (destino_key, costo_compartido, costo_privado, comision)
values
%s
on conflict (destino_key) do update set
  costo_compartido = excluded.costo_compartido,
  costo_privado = excluded.costo_privado,
  comision = excluded.comision,
  actualizado_at = now();

-- Verificacion: si algo quedo en null, ese destino quedo sin costo.
select destino_key, nombre, costo_compartido, costo_privado, comision,
       precio_compartido, precio_privado
  from public.transfer_destinos
 where costo_compartido is not null or costo_privado is not null
 order by destino_key;
""" % valores

with open(OUT, "w", encoding="utf-8") as f:
    f.write(sql)

print("OK ->", OUT)
print("filas con costo:", len(filas))
print("completas (las dos modalidades + comision):",
      sum(1 for _, _, c, p, m in filas if c is not None and p is not None and m is not None))
print("a medio cargar:", sum(1 for _, _, c, p, m in filas if not (c is not None and p is not None and m is not None)))
