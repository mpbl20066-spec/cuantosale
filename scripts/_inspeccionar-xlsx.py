# -*- coding: utf-8 -*-
"""Muestra la estructura de un xlsx: hojas, columnas y las primeras filas.

    python scripts/inspeccionar-xlsx.py data/tours.xlsx
    python scripts/inspeccionar-xlsx.py data/tours.xlsx --filas 5

Es un script de lectura, no escribe nada. Sirve para ver que trae un archivo
que todavia no se conoce, sin abrirlo a mano ni arriesgarse a subirlo con
encabezados que no son los que espera el importador.
"""
import os, sys
from openpyxl import load_workbook

if len(sys.argv) < 2:
    sys.exit(__doc__)

RUTA = sys.argv[1]
FILAS = 5
if "--filas" in sys.argv:
    FILAS = int(sys.argv[sys.argv.index("--filas") + 1])

if not os.path.exists(RUTA):
    sys.exit("No existe: " + RUTA)

wb = load_workbook(RUTA, data_only=True, read_only=True)
print("archivo:", RUTA)
print("hojas:", wb.sheetnames)
print("")

for nombre in wb.sheetnames:
    ws = wb[nombre]
    filas = ws.iter_rows(values_only=True)
    try:
        head = next(filas)
    except StopIteration:
        print("[%s] vacia" % nombre)
        print("")
        continue

    # Cuantas filas hay, sin cargar todo en memoria.
    total = ws.max_row or 0
    print("[%s]  %s filas x %s columnas" % (nombre, total, len(head)))
    print("  columnas:")
    for i, h in enumerate(head):
        etiqueta = (str(h).strip() if h is not None else "")
        print("    %2d  %s" % (i + 1, etiqueta if etiqueta else "(sin encabezado)"))
    print("  primeras filas:")
    for n, fila in enumerate(filas):
        if n >= FILAS:
            break
        celdas = []
        for v in fila:
            s = "" if v is None else str(v).strip()
            if len(s) > 46:
                s = s[:43] + "..."
            celdas.append(s)
        print("    %d | %s" % (n + 2, " | ".join(celdas)))
    print("")
