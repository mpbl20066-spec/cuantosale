# -*- coding: utf-8 -*-
"""Exporta data/tours.json a data/tours-editar.xlsx para editar precios y textos.

    python scripts/exportar-tours-xlsx.py

Para volver: python scripts/importar-tours-xlsx.py (ver ese archivo).
"""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = os.path.join(os.path.dirname(__file__), '..')
d = json.load(open(os.path.join(RAIZ, 'data', 'tours.json'), encoding='utf8'))
meta = d['_meta']

wb = Workbook()
ws = wb.active
ws.title = 'Tours'
F = 'Arial'
cols = [('Código destino', 10), ('Destino', 24), ('Título', 42), ('Descripción', 50),
        ('Detalle', 60), ('Precio R$ (origen)', 14), ('Precio USD', 12)]
for i, (n, w) in enumerate(cols, 1):
    c = ws.cell(row=1, column=i, value=n)
    c.font = Font(name=F, bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor='1F3864')
    c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.column_dimensions[c.column_letter].width = w
ws.row_dimensions[1].height = 32

azul = Font(name=F, color='0000FF')
negro = Font(name=F)
gris = PatternFill('solid', fgColor='EDEDED')
for r, t in enumerate(d['tours'], 2):
    vals = [t['destinos'][0], t['destino'], t['titulo'], t['descripcion'], t['detalle'],
            t.get('precio_brl'), t['precio']]
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = azul if i >= 4 else negro
        c.alignment = Alignment(wrap_text=i in (4, 5), vertical='top')
        if i <= 3:
            c.fill = gris  # clave de la fila: no se edita
        if i >= 6:
            c.number_format = '#,##0.00'
ws.freeze_panes = 'D2'
ws.auto_filter.ref = f'A1:G{len(d["tours"]) + 1}'

L = wb.create_sheet('Leeme')
L.column_dimensions['A'].width = 110
lineas = [
    'CÓMO USAR ESTA PLANILLA',
    '',
    'Editá solo las columnas en azul: Descripción, Detalle, Precio R$ y Precio USD.',
    'NO cambies Código destino, Destino ni Título (columnas grises): son la clave con la que se identifica cada tour en Supabase.',
    'No borres ni agregues filas: el importador solo actualiza tours que ya existen.',
    '',
    'PRECIOS',
    f'- Si el tour tiene Precio R$, manda ese: el USD se recalcula como R$ / {meta["cotizacion_brl_usd"]} + {meta["margen_usd"]} de margen (los valores salen de data/tours.json, _meta).',
    '- Si Precio R$ está vacío, el Precio USD que escribas es el que se usa (ya con margen incluido).',
    '- Para pasar un tour de reales a dólares, vaciá Precio R$ y escribí el Precio USD.',
    '',
    'PARA APLICAR LOS CAMBIOS (guardá el xlsx y cerralo antes):',
    '  python scripts/importar-tours-xlsx.py --dry-run     (muestra qué cambiaría)',
    '  python scripts/importar-tours-xlsx.py               (escribe data/tours.json)',
    '  npm run build:tours                                 (regenera la web)',
    '  node scripts/cargar-tours.js                        (sube a Supabase)',
]
for i, t in enumerate(lineas, 1):
    c = L.cell(row=i, column=1, value=t)
    c.font = Font(name=F, bold=(i in (1, 7)))
out = os.path.join(RAIZ, 'data', 'tours-editar.xlsx')
wb.save(out)
print('ok', len(d['tours']), 'tours ->', out)
